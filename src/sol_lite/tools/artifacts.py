"""Portable, permission-checked artifact authoring. No Microsoft runtime required."""

from __future__ import annotations

import importlib.util
import json
import os
import posixpath
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from pydantic import BaseModel, ConfigDict, Field

from ..capabilities.risk import RiskClass
from ..permissions.scopes import READ, WRITE
from ..security.path_guard import resolve_workspace_path
from .base import ToolDefinition


class SpecModel(BaseModel):
    model_config = ConfigDict(extra="forbid")


class Theme(SpecModel):
    font: str = "Arial"
    accent: str = Field(default="2457A7", pattern=r"^[0-9A-Fa-f]{6}$")
    text: str = Field(default="202938", pattern=r"^[0-9A-Fa-f]{6}$")


class Section(SpecModel):
    heading: str
    paragraphs: list[str] = Field(default_factory=list)
    bullets: list[str] = Field(default_factory=list)
    table: list[list[str]] = Field(default_factory=list)


class Slide(SpecModel):
    title: str
    bullets: list[str] = Field(default_factory=list, max_length=6)
    notes: str = ""


class Chart(SpecModel):
    title: str
    first_row: int = Field(ge=1)
    last_row: int = Field(ge=2)
    category_column: int = Field(default=1, ge=1)
    value_column: int = Field(default=2, ge=1)
    anchor: str | None = Field(default=None, pattern=r"^[A-Z]{1,3}[1-9][0-9]{0,5}$")


class Sheet(SpecModel):
    name: str = Field(min_length=1, max_length=31)
    rows: list[list[str | int | float | bool | None]]
    charts: list[Chart] = Field(default_factory=list)


class ArtifactSpec(SpecModel):
    title: str = Field(min_length=1, max_length=200)
    sections: list[Section] = Field(default_factory=list)
    slides: list[Slide] = Field(default_factory=list)
    sheets: list[Sheet] = Field(default_factory=list)
    mermaid: str = Field(default="", max_length=100000)
    theme: Theme = Field(default_factory=Theme)


def discover_dependencies():
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if not office:
        candidates = [Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")]
        for key in ("PROGRAMFILES", "PROGRAMFILES(X86)"):
            if os.environ.get(key):
                candidates.append(Path(os.environ[key]) / "LibreOffice/program/soffice.exe")
        office = next((str(p) for p in candidates if p.is_file()), None)
    # Avoid .cmd shell execution on Windows: invoke the installed JS entrypoint via Node.
    mermaid = shutil.which("mmdc")
    node = shutil.which("node")
    command = [mermaid] if mermaid else None
    if mermaid and Path(mermaid).suffix.lower() in {".cmd", ".bat", ".ps1"}:
        entry = Path(mermaid).parent / "node_modules/@mermaid-js/mermaid-cli/src/cli.js"
        command = [node, str(entry)] if node and entry.is_file() else None
    return {
        "python": {
            name: importlib.util.find_spec(name) is not None
            for name in ("docx", "pptx", "openpyxl", "fitz")
        },
        "libreoffice": office,
        "mermaid_command": command,
        "setup": "Install sol-lite[artifacts], LibreOffice, and @mermaid-js/mermaid-cli locally. No Microsoft Office required.",
    }


def artifact_dependencies(ctx, args):
    ctx.permission_engine.check(READ)
    return discover_dependencies()


def _run(command):
    try:
        result = subprocess.run(
            command, capture_output=True, text=True, timeout=60, shell=False, check=False
        )
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError("Artifact renderer exceeded 60 seconds.") from exc
    if result.returncode:
        raise RuntimeError(f"Artifact renderer failed: {result.stderr[-2000:]}")


def _check_mermaid(source):
    if not source.strip():
        raise ValueError("Mermaid source is required.")
    if re.search(r"%%\s*\{|\b(click|href|callback)\b|<|https?://|file:", source, re.IGNORECASE):
        raise ValueError(
            "Mermaid directives, links, callbacks, HTML, and external resources are unsupported."
        )


def _author(path, spec):
    suffix = path.suffix.lower()
    allowed = {".docx": "sections", ".pptx": "slides", ".xlsx": "sheets", ".mmd": "mermaid"}
    for field in ("sections", "slides", "sheets", "mermaid"):
        if getattr(spec, field) and field != allowed.get(suffix):
            raise ValueError(f"Content field {field} does not belong in {suffix} artifacts.")
    if suffix == ".docx":
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Mm, Pt, RGBColor

        doc = Document()
        section = doc.sections[0]
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = Mm(22)
        section.left_margin = section.right_margin = Mm(24)
        normal = doc.styles["Normal"]
        normal.font.name, normal.font.size = spec.theme.font, Pt(11)
        normal.font.color.rgb = RGBColor.from_string(spec.theme.text)
        normal.paragraph_format.space_after = Pt(8)
        normal.paragraph_format.line_spacing = 1.15
        title_style = doc.styles["Title"]
        title_style.font.name, title_style.font.size = spec.theme.font, Pt(28)
        title_style.font.color.rgb = RGBColor.from_string("000000")
        for name in ("Heading 1", "Heading 2"):
            doc.styles[name].font.name = spec.theme.font
            doc.styles[name].font.color.rgb = RGBColor.from_string(spec.theme.accent)
            doc.styles[name].paragraph_format.keep_with_next = True
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run("Page ").font.size = Pt(9)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
        doc.add_heading(spec.title, 0)
        for item in spec.sections:
            doc.add_heading(item.heading, 1)
            for value in item.paragraphs:
                doc.add_paragraph(value)
            for value in item.bullets:
                doc.add_paragraph(value, "List Bullet")
            if item.table:
                width = len(item.table[0])
                if not width or any(len(row) != width for row in item.table):
                    raise ValueError("Document tables must be rectangular.")
                table = doc.add_table(rows=0, cols=width)
                table.style = "Light Shading Accent 1"
                for row in item.table:
                    for cell, value in zip(table.add_row().cells, row):
                        cell.text = value
                header = table.rows[0]
                repeated = OxmlElement("w:tblHeader")
                header._tr.get_or_add_trPr().append(repeated)
                for cell in header.cells:
                    shading = OxmlElement("w:shd")
                    shading.set(qn("w:fill"), spec.theme.accent)
                    cell._tc.get_or_add_tcPr().append(shading)
                    for run in cell.paragraphs[0].runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor.from_string("FFFFFF")
        doc.save(path)
    elif suffix == ".pptx":
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Inches, Pt

        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
        for number, item in enumerate(spec.slides or [Slide(title=spec.title)], 1):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            for top, height, values, size, colour in (
                (0.6, 1.1, [item.title], 30, spec.theme.accent),
                (2.0, 4.8, item.bullets, 22, spec.theme.text),
            ):
                frame = slide.shapes.add_textbox(
                    Inches(0.8), Inches(top), Inches(11.7), Inches(height)
                ).text_frame
                frame.word_wrap = True
                for index, value in enumerate(values):
                    p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
                    p.text = ("• " if top == 2.0 else "") + value
                    p.font.name, p.font.size = spec.theme.font, Pt(size)
                    p.font.color.rgb = RGBColor.from_string(colour)
                    p.space_after = Pt(18)
            footer = slide.shapes.add_textbox(
                Inches(0.8), Inches(7.0), Inches(11.7), Inches(0.3)
            ).text_frame
            footer.paragraphs[0].text = f"{spec.title}  |  {number}"
            footer.paragraphs[0].font.name = spec.theme.font
            footer.paragraphs[0].font.size = Pt(10)
            footer.paragraphs[0].font.color.rgb = RGBColor.from_string(spec.theme.text)
            slide.notes_slide.notes_text_frame.text = item.notes
        deck.save(path)
    elif suffix == ".xlsx":
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

        book = Workbook()
        book.remove(book.active)
        if not spec.sheets:
            raise ValueError("At least one sheet is required.")
        for item in spec.sheets:
            if re.search(r"[\\/*?:\[\]]", item.name) or item.name in book.sheetnames:
                raise ValueError("Sheet names must be valid and unique.")
            sheet = book.create_sheet(item.name)
            for row in item.rows:
                for value in row:
                    if (
                        isinstance(value, str)
                        and value.startswith("=")
                        and re.search(
                            r"\[|https?://|\b(HYPERLINK|WEBSERVICE|RTD|DDE)\b|\|",
                            value,
                            re.IGNORECASE,
                        )
                    ):
                        raise ValueError("External or active workbook formulas are unsupported.")
                sheet.append(row)
            for row in sheet:
                for cell in row:
                    cell.font = Font(name=spec.theme.font, size=11, color=spec.theme.text)
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor=spec.theme.accent)
                cell.font = Font(name=spec.theme.font, bold=True, color="FFFFFF")
            for index in range(1, sheet.max_column + 1):
                sheet.column_dimensions[get_column_letter(index)].width = 22
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.print_options.horizontalCentered = True
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.orientation = "landscape"
            sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
            sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
            print_last_row, print_last_column = sheet.max_row, sheet.max_column
            for chart_spec in item.charts:
                if (
                    chart_spec.last_row > sheet.max_row
                    or chart_spec.first_row >= chart_spec.last_row
                ):
                    raise ValueError("Chart row range is invalid.")
                if max(chart_spec.category_column, chart_spec.value_column) > sheet.max_column:
                    raise ValueError("Chart column range is invalid.")
                chart = BarChart()
                chart.title = chart_spec.title
                chart.style = 10
                chart.add_data(
                    Reference(
                        sheet,
                        min_col=chart_spec.value_column,
                        min_row=chart_spec.first_row,
                        max_row=chart_spec.last_row,
                    ),
                    titles_from_data=True,
                )
                chart.set_categories(
                    Reference(
                        sheet,
                        min_col=chart_spec.category_column,
                        min_row=chart_spec.first_row + 1,
                        max_row=chart_spec.last_row,
                    )
                )
                chart.series[0].graphicalProperties.solidFill = spec.theme.accent
                anchor = chart_spec.anchor or f"A{sheet.max_row + 3}"
                column, row = coordinate_from_string(anchor)
                print_last_row = max(print_last_row, row + 16)
                print_last_column = max(print_last_column, column_index_from_string(column) + 4)
                sheet.add_chart(chart, anchor)
            sheet.print_area = f"A1:{get_column_letter(print_last_column)}{print_last_row}"
        book.save(path)
    elif suffix == ".mmd":
        _check_mermaid(spec.mermaid)
        path.write_text(spec.mermaid, encoding="utf-8")
    else:
        raise ValueError(
            "Authoring supports .docx, .pptx, .xlsx, and .mmd; request PDF through rendering."
        )


def _validate(path):
    try:
        return _validate_content(path)
    except ImportError as exc:
        raise RuntimeError("Missing validation dependency. Install sol-lite[artifacts].") from exc
    except (zipfile.BadZipFile, ElementTree.ParseError, KeyError) as exc:
        raise ValueError(f"Invalid artifact package: {exc}") from exc


def _validate_content(path):
    suffix = path.suffix.lower()
    warnings = []
    if suffix == ".mmd":
        _check_mermaid(path.read_text(encoding="utf-8"))
        warnings.append("Mermaid syntax requires successful renderer execution.")
    elif suffix in {".docx", ".pptx", ".xlsx"}:
        with zipfile.ZipFile(path) as archive:
            if sum(info.file_size for info in archive.infolist()) > 50_000_000:
                raise ValueError("Office package exceeds the 50 MB inspection limit.")
            if archive.testzip() or "[Content_Types].xml" not in archive.namelist():
                raise ValueError("Invalid Office package.")
            for name in archive.namelist():
                if name.endswith(".rels"):
                    root = ElementTree.fromstring(archive.read(name))
                    if any(rel.attrib.get("TargetMode") == "External" for rel in root):
                        raise ValueError("External Office relationships cannot be rendered safely.")
                    base = posixpath.dirname(posixpath.dirname(name))
                    for rel in root:
                        target = rel.attrib.get("Target", "")
                        part = posixpath.normpath(posixpath.join(base, target)).lstrip("/")
                        if part not in archive.namelist():
                            raise ValueError(f"Office package references missing asset: {part}")
                if "vbaproject" in name.lower() or "/embeddings/" in name.lower():
                    raise ValueError("Macros and embedded executable objects are unsupported.")
        if suffix == ".docx":
            from docx import Document

            document = Document(path)
            if any(len(table.columns) > 7 for table in document.tables):
                warnings.append("Wide document table; inspect rendered column widths.")
        if suffix == ".xlsx":
            from openpyxl import load_workbook

            book = load_workbook(path, data_only=False)
            for sheet in book:
                for row in sheet:
                    if any(cell.data_type == "e" for cell in row):
                        warnings.append(f"Formula/cell error on {sheet.title}.")
                    for cell in row:
                        if cell.data_type == "f" and re.search(
                            r"\[|https?://|\b(HYPERLINK|WEBSERVICE|RTD|DDE)\b|\|",
                            cell.value,
                            re.IGNORECASE,
                        ):
                            raise ValueError(
                                "External or active workbook formulas are unsupported."
                            )
            book.close()
            warnings.append(
                "Formula values require spreadsheet recalculation; cached values are not proof of correctness."
            )
        if suffix == ".pptx":
            from pptx import Presentation

            for i, slide in enumerate(Presentation(path).slides, 1):
                if any(len(shape.text) > 700 for shape in slide.shapes if shape.has_text_frame):
                    warnings.append(f"Slide {i} has dense text; inspect preview for overflow.")
    else:
        raise ValueError("Unsupported artifact type.")
    return {"structural": "passed", "warnings": warnings, "visual_review": "not_performed"}


def _render(source, directory, dependencies, theme=None):
    outputs, previews, warnings = [], [], []
    if source.suffix.lower() == ".mmd":
        if not dependencies["mermaid_command"]:
            return outputs, previews, ["Mermaid renderer missing; diagram exports unverified."]
        theme = theme or Theme()
        config = directory / "mermaid-config.json"
        config.write_text(
            json.dumps(
                {
                    "securityLevel": "strict",
                    "theme": "base",
                    "themeVariables": {
                        "fontFamily": theme.font,
                        "primaryColor": "#EEF3FA",
                        "primaryTextColor": "#" + theme.text,
                        "primaryBorderColor": "#" + theme.accent,
                        "lineColor": "#" + theme.accent,
                    },
                }
            ),
            encoding="utf-8",
        )
        for extension in ("svg", "png"):
            output = directory / f"{source.stem}.{extension}"
            _run(
                [
                    *dependencies["mermaid_command"],
                    "-i",
                    str(source),
                    "-o",
                    str(output),
                    "-b",
                    "white",
                    "-s",
                    "2",
                    "-c",
                    str(config),
                ]
            )
            if not output.is_file() or not output.stat().st_size:
                raise RuntimeError("Mermaid renderer did not produce the requested output.")
            outputs.append(output)
            if extension == "png":
                previews.append(output)
    else:
        if not dependencies["libreoffice"]:
            return outputs, previews, ["LibreOffice missing; PDF and previews unavailable."]
        profile = directory / "office-profile"
        _run(
            [
                dependencies["libreoffice"],
                f"-env:UserInstallation={profile.as_uri()}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(directory),
                str(source),
            ]
        )
        pdf = directory / f"{source.stem}.pdf"
        if not pdf.is_file() or not pdf.stat().st_size:
            raise RuntimeError("LibreOffice did not produce a PDF.")
        outputs.append(pdf)
        if dependencies["python"]["fitz"]:
            import fitz

            with fitz.open(pdf) as document:
                if len(document) > 200:
                    raise ValueError("Preview limit is 200 pages; split the artifact.")
                for index, page in enumerate(document):
                    preview = directory / f"{source.stem}-page-{index + 1}.png"
                    page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(preview)
                    previews.append(preview)
        else:
            warnings.append("PyMuPDF missing; PDF produced but page previews unavailable.")
    return outputs, previews, warnings


def _publish(files, destination, workspace):
    # Exclusive creation prevents silent overwrite, including races.
    targets = [resolve_workspace_path(destination / p.name, workspace) for p in files]
    if any(p.exists() for p in targets):
        raise FileExistsError(
            "Artifact output exists; choose a fresh output directory or filename."
        )
    destination.mkdir(parents=True, exist_ok=True)
    created = []
    try:
        for source, target in zip(files, targets):
            with target.open("xb") as stream:
                created.append(target)
                with source.open("rb") as incoming:
                    shutil.copyfileobj(incoming, stream)
    except Exception:
        for target in created:
            target.unlink(missing_ok=True)
        raise
    return [str(p) for p in targets]


def _approve(ctx, args, operation, target):
    bound_arguments = {key: value for key, value in args.items() if key != "approval_id"}
    ctx.permission_engine.check(
        WRITE,
        approval_id=args.get("approval_id"),
        operation=operation,
        target=str(target),
        arguments=bound_arguments,
        plan=args.get("plan", ""),
    )


def artifact_create(ctx, args):
    spec = ArtifactSpec.model_validate(args["spec"])
    target = resolve_workspace_path(args["path"], ctx.workspace_root)
    if target.exists():
        raise FileExistsError(target)
    _approve(ctx, args, "artifact_create", target)
    with tempfile.TemporaryDirectory(prefix="sol-artifact-") as temp:
        source = Path(temp) / target.name
        try:
            _author(source, spec)
            validation = _validate(source)
        except ImportError as exc:
            raise RuntimeError(
                "Missing authoring dependency. Install sol-lite[artifacts]."
            ) from exc
        files, previews, warnings = [], [], []
        if args.get("render", True):
            try:
                files, previews, warnings = _render(
                    source, Path(temp), discover_dependencies(), spec.theme
                )
            except (RuntimeError, OSError, ValueError) as exc:
                warnings.append(str(exc))
        validation["warnings"].extend(warnings)
        outputs = _publish(
            [source, *files, *[p for p in previews if p not in files]],
            target.parent,
            ctx.workspace_root,
        )
    result = {
        "ok": True,
        "status": "rendered_unreviewed" if previews else "authored_unverified",
        "outputs": outputs,
        "previews": [str(target.parent / p.name) for p in previews],
        "validation": validation,
    }
    ctx.audit.record("artifact.create", path=str(target), status=result["status"])
    return result


def artifact_validate(ctx, args):
    ctx.permission_engine.check(READ)
    source = resolve_workspace_path(args["path"], ctx.workspace_root, must_exist=True)
    return _validate(source)


def artifact_render(ctx, args):
    ctx.permission_engine.check(READ)
    source = resolve_workspace_path(args["path"], ctx.workspace_root, must_exist=True)
    destination = resolve_workspace_path(args["output_directory"], ctx.workspace_root)
    _validate(source)
    _approve(ctx, args, "artifact_render", destination)
    with tempfile.TemporaryDirectory(prefix="sol-render-") as temp:
        # Render a private copy, never the user's editable original.
        staged = Path(temp) / source.name
        shutil.copyfile(source, staged)
        files, previews, warnings = _render(staged, Path(temp), discover_dependencies())
        outputs = (
            _publish(
                [*files, *[p for p in previews if p not in files]], destination, ctx.workspace_root
            )
            if files
            else []
        )
    return {
        "ok": bool(outputs),
        "outputs": outputs,
        "previews": [str(destination / p.name) for p in previews],
        "warnings": warnings,
        "visual_review": "not_performed",
    }


def artifact_tools():
    common = {"approval_id": {"type": "string"}, "plan": {"type": "string"}}
    spec_schema = ArtifactSpec.model_json_schema()
    definitions_schema = spec_schema.pop("$defs", {})
    definitions = [
        ("artifact_dependencies", artifact_dependencies, {}, [], False),
        (
            "artifact_create",
            artifact_create,
            {
                **common,
                "path": {"type": "string"},
                "spec": spec_schema,
                "render": {"type": "boolean", "default": True},
            },
            ["path", "spec"],
            True,
        ),
        (
            "artifact_render",
            artifact_render,
            {**common, "path": {"type": "string"}, "output_directory": {"type": "string"}},
            ["path", "output_directory"],
            True,
        ),
        ("artifact_validate", artifact_validate, {"path": {"type": "string"}}, ["path"], False),
    ]
    return [
        ToolDefinition(
            name,
            {
                "artifact_dependencies": "Check local artifact dependencies without installing anything.",
                "artifact_create": "Create premium editable DOCX/PPTX/XLSX or Mermaid source from typed content; optionally export PDF or SVG/PNG. Visual review remains separate.",
                "artifact_render": "Export an existing artifact to PDF or diagram SVG/PNG and page previews without changing the original.",
                "artifact_validate": "Check artifact structure; does not certify visual quality.",
            }[name],
            {
                "type": "object",
                "properties": properties,
                "required": required,
                "additionalProperties": False,
                **({"$defs": definitions_schema} if name == "artifact_create" else {}),
            },
            handler,
            permissions=(WRITE,) if mutating else (READ,),
            risk=(RiskClass.MUTATING,) if mutating else (RiskClass.READ_ONLY,),
            mutability=mutating,
        )
        for name, handler, properties, required, mutating in definitions
    ]
